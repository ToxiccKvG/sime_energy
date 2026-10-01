"""Génère le classeur Excel et le rapport PDF à partir de tables.pkl / clean.pkl / figs/."""
import pickle, os, shutil
import numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether

T = pickle.load(open("tables.pkl", "rb")); clean = pd.read_pickle("clean.pkl"); flags = pd.read_pickle("flags.pkl")
raw = pd.read_pickle("raw.pkl"); raw.columns = ["ts"] + list(raw.columns[1:])
OUT = "out"; os.makedirs(OUT, exist_ok=True); DESK = os.path.expanduser("~/Desktop/Meteo_Donsin_2026-09")
D, H, S, prof, ind, log = T["D"], T["H"], T["S"], T["prof"], T["ind"], T["log"]
NAVY = "1F3A5F"

# ======================= EXCEL =======================
cols_fr = {"temperature": "Température (°C)", "humidity": "Humidité (%)", "uv_index": "Indice UV", "illuminance_lux": "Éclairement (lux)",
           "wind_speed_ms": "Vent moyen (m/s)", "wind_gust_ms": "Rafale (m/s)", "wind_direction_deg": "Direction vent (°)",
           "pressure_hpa": "Pression (hPa)", "precipitation_mm": "Pluie cumul compteur (mm)", "dewpoint_c": "Point de rosée (°C)",
           "feels_like_c": "Ressenti (°C)", "battery_voltage_v": "Tension batterie (V)", "signal_rssi": "Signal Wi-Fi (dBm)",
           "moisture_alarm": "Alarme humidité boîtier", "flag_anomalie": "Anomalie détectée", "jour_degrade": "Jour dégradé"}
cl = clean.reset_index().rename(columns={"measured_at": "Horodatage mesure (UTC)", **cols_fr})
readme = pd.DataFrame({"Contenu": [
    "Station météo Donsin — Septembre 2026 : données nettoyées et analysées",
    "Appareil : Shelly BLU / station SBWS-90CM (device_id XB211299189491624), site Donsin, salle CENTRALE METEO_DONSIN",
    "Source : table brute shelly_cl (Supabase), extraite le 01/10/2026. Horodatages en UTC (= heure locale, UTC+0).",
    "",
    "ONGLETS",
    "Synthèse : indicateurs clés et statistiques descriptives",
    "Données_propres : 1 ligne par mesure distincte (clé measured_at), colonnes vides retirées, pics/aberrations exclus (NaN)",
    "Horaire : moyennes horaires ; 'heure_valide' = au moins 30 mesures sur 60",
    "Journalier : agrégats par jour ; 'jour_complet' = au moins 8 h valides (la station ne transmet que ~10 h/jour)",
    "Profil_horaire : cycle journalier moyen (par heure UTC)",
    "Rose_vents : fréquence (%) par secteur de 22,5° et classe de vitesse",
    "Pluie : événements d'incrément du compteur de pluie",
    "Trous_transmission : silences > 10 min entre deux mesures",
    "Correlations : corrélations de Pearson sur les moyennes horaires",
    "Journal_nettoyage : chaque règle appliquée et le nombre de lignes concernées",
    "Graphiques : figures du rapport",
    "",
    "POINTS D'ATTENTION",
    "1. 62 % des 43 200 lignes brutes sont des doublons de collecte (même measured_at) : le poller relit la dernière valeur connue chaque minute.",
    "2. La station ne transmet pas la nuit (≈ 19 h → 07 h UTC) : aucun minimum nocturne fiable, aucun 'nuit tropicale' calculable.",
    "3. Le cumul de pluie est un compteur : la pluie des nuits muettes apparaît d'un coup au réveil de la station.",
    "4. 3 jours quasi vides (20, 22, 23/09) : station noyée / hors ligne (alarme d'humidité du boîtier active).",
]})
with pd.ExcelWriter(f"{OUT}/Meteo_Donsin_2026-09_analyse.xlsx", engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm:ss") as xw:
    readme.to_excel(xw, sheet_name="Lisez-moi", index=False)
    syn = pd.DataFrame([(k, v) for k, v in ind.items()], columns=["Indicateur", "Valeur"])
    syn.to_excel(xw, sheet_name="Synthèse", index=False); S.to_excel(xw, sheet_name="Synthèse", index=False, startrow=len(syn) + 3)
    cl.to_excel(xw, sheet_name="Données_propres", index=False)
    Hx = H.copy(); Hx.index.name = "Heure (UTC)"; Hx.reset_index().to_excel(xw, sheet_name="Horaire", index=False)
    Dx = D.copy(); Dx.index.name = "Jour"; Dx.reset_index().to_excel(xw, sheet_name="Journalier", index=False)
    prof.reset_index().rename(columns={"measured_at": "Heure UTC", "index": "Heure UTC"}).to_excel(xw, sheet_name="Profil_horaire", index=False)
    T["rose_pct"].reset_index().rename(columns={"index": "Secteur"}).to_excel(xw, sheet_name="Rose_vents", index=False)
    T["rain_events"].reset_index().rename(columns={"measured_at": "Horodatage (UTC)", "mm": "Incrément (mm)"}).to_excel(xw, sheet_name="Pluie", index=False)
    g = T["gap_tbl"].copy(); g["debut"] = g.debut.dt.round("s"); g["fin"] = g.fin.dt.round("s"); g.to_excel(xw, sheet_name="Trous_transmission", index=False)
    T["corr"].to_excel(xw, sheet_name="Correlations")
    pd.DataFrame(log, columns=["Étape", "Règle / contrôle", "Résultat"]).astype({"Résultat": str}).to_excel(xw, sheet_name="Journal_nettoyage", index=False)
    pd.DataFrame({"Graphiques": ["Voir ci-dessous"]}).to_excel(xw, sheet_name="Graphiques", index=False)

wb = load_workbook(f"{OUT}/Meteo_Donsin_2026-09_analyse.xlsx")
hdr_fill = PatternFill("solid", fgColor=NAVY)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.row == 1 or (ws.title == "Synthèse" and c.value in ("Variable",)):
                pass
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2" if ws.title in ("Données_propres", "Horaire", "Journalier") else "A2"
    for i, col in enumerate(ws.columns, 1):
        w = max((len(str(c.value)) for c in list(col)[:300] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 70 if ws.title in ("Lisez-moi", "Journal_nettoyage", "Synthèse") else 24)
ws = wb["Synthèse"]
for r in ws.iter_rows():
    if r[0].value == "Variable":
        for c in r: c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
wb["Lisez-moi"]["A1"].font = Font(bold=True, size=13, color=NAVY)
wg = wb["Graphiques"]; rowp = 3
for f in sorted(os.listdir("figs")):
    im = XLImage(f"figs/{f}"); im.width, im.height = im.width * 0.55, im.height * 0.55
    wg.add_image(im, f"A{rowp}"); rowp += int(im.height / 20) + 3
wb.save(f"{OUT}/Meteo_Donsin_2026-09_analyse.xlsx")

# brut corrigé (en-têtes propres)
raw.to_excel(f"{OUT}/Meteo_Donsin_2026-09_brut.xlsx", index=False, sheet_name="brut shelly_cl")
if os.path.exists(os.path.expanduser("~/Desktop/Meteo_Donsin_septembre_2026.xlsx")):
    os.remove(os.path.expanduser("~/Desktop/Meteo_Donsin_septembre_2026.xlsx"))

# ======================= PDF =======================
ss = getSampleStyleSheet()
st = {"t": ParagraphStyle("t", parent=ss["Title"], textColor=colors.HexColor("#" + NAVY), fontSize=22, leading=26, alignment=0),
      "h1": ParagraphStyle("h1", parent=ss["Heading1"], textColor=colors.HexColor("#" + NAVY), fontSize=14, spaceBefore=10, spaceAfter=6),
      "h2": ParagraphStyle("h2", parent=ss["Heading2"], textColor=colors.HexColor("#" + NAVY), fontSize=11, spaceBefore=8, spaceAfter=3),
      "b": ParagraphStyle("b", parent=ss["BodyText"], fontSize=9.5, leading=13.5),
      "s": ParagraphStyle("s", parent=ss["BodyText"], fontSize=8, leading=10.5, textColor=colors.HexColor("#555555")),
      "k": ParagraphStyle("k", parent=ss["BodyText"], fontSize=17, leading=20, alignment=1, textColor=colors.HexColor("#" + NAVY), fontName="Helvetica-Bold"),
      "kl": ParagraphStyle("kl", parent=ss["BodyText"], fontSize=7.5, leading=9, alignment=1, textColor=colors.HexColor("#555555"))}
P = lambda t, s="b": Paragraph(t, st[s])
def kpis(items):
    cells = [[P(v, "k") for v, _ in items], [P(l, "kl") for _, l in items]]
    t = Table(cells, colWidths=[17 * cm / len(items)] * len(items))
    t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .5, colors.HexColor("#c9d3e0")), ("INNERGRID", (0, 0), (-1, -1), .25, colors.HexColor("#e3e8ef")),
                           ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f4f7fb")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return t
def tbl(df, widths=None, fs=7.5):
    data = [list(df.columns)] + df.astype(str).values.tolist()
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), fs),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7fb")]), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#d5dbe5")), ("ALIGN", (1, 1), (-1, -1), "RIGHT")]))
    return t
def fig(name, w=17):
    im = Image(f"figs/{name}.png"); r = im.imageHeight / im.imageWidth; im.drawWidth = w * cm; im.drawHeight = w * cm * r; return im

Df = D[D.jour_complet]
dfull = int(T["nfull"])
Tmax_d = D.T_max.idxmax(); Tmin_d = D.T_min.idxmin()
uvdays = int((D.UV_max >= 8).sum())
p_peak = prof.pressure_hpa.idxmax(), prof.pressure_hpa.idxmin()
story = [P("Station météo Donsin", "t"), P("Nettoyage et analyse des données — septembre 2026", "h1"),
         P("Source : table shelly_cl (Supabase) · Appareil SBWS-90CM <i>XB211299189491624</i> · Période 01/09 → 30/09/2026 · Horodatages UTC (= heure locale)", "s"), Spacer(1, .4 * cm)]
story += [P("1. À retenir", "h1"),
          kpis([("38,5 %", "des 43 200 lignes brutes sont des mesures distinctes"), ("07:20 → 17:56", "fenêtre de transmission médiane (UTC)"),
                (f"{dfull}/30", "jours exploitables (≥ 8 h valides)"), ("58,6 mm", "pluie enregistrée sur le mois")]), Spacer(1, .3 * cm),
          P("<b>La station météo fonctionne, mais ne transmet que ~10 h par jour.</b> Entre 19 h et 07 h UTC, aucune mesure n'arrive "
            f"({ind['Heures valides entre 06 h et 19 h UTC (seules heures observées)']} heures valides sur 720, toutes en journée). "
            "Les statistiques ci-dessous décrivent donc le <b>jour</b> ; les minima nocturnes, les « nuits tropicales » et les moyennes sur 24 h ne sont pas calculables."),
          Spacer(1, .15 * cm),
          P("<b>Sur la fenêtre observée</b>, septembre est chaud et humide : température moyenne 31,1 °C (pic 37,2 °C le "
            f"{Tmax_d:%d/%m}), humidité 64 % en moyenne (de 54 % à 14 h jusqu'à 90 % à 06 h), point de rosée stable à 23,2 °C, ressenti jusqu'à 44,3 °C. "
            f"L'indice UV atteint 8 ou plus sur {uvdays} jours."),
          Spacer(1, .15 * cm),
          P(f"<b>Pluie : {ind['Pluie enregistrée sur le mois (mm)']} mm</b> sur 11 jours, mais {ind['Pluie tombée pendant les nuits muettes (mm)']} mm (89 %) apparaissent d'un coup au "
            "réveil de la station : la pluie nocturne est comptée, mais impossible à dater. Seulement "
            f"{ind['Pluie observée en direct (mm)']} mm ont été vus en direct.")]
story += [P("2. Qualité des données et nettoyage", "h1"),
          P("<b>Doublons de collecte.</b> Le poller lit la dernière valeur connue chaque minute ; quand la station ne transmet pas, il stocke la même mesure "
            "(même <i>measured_at</i>) en boucle. Sur 43 200 lignes, <b>16 632 mesures distinctes</b> : 61,5 % de lignes redondantes, avec un retard "
            "collecte→mesure médian de 2 h 27 et jusqu'à 32 h 30. Tout est dédoublonné sur <i>measured_at</i> avant analyse."),
          P("<b>Plausibilité.</b> Aucune valeur hors bornes physiques, aucune incohérence (rafale &lt; vent moyen, rosée &gt; température), aucune séquence figée "
            "≥ 120 mesures, un seul pic isolé de vent écarté (filtre de Hampel). Le capteur est propre ; le problème est la <b>disponibilité</b>, pas la justesse."),
          P("<b>Jours dégradés.</b> Les 20, 22 et 23/09 comptent moins de 100 mesures, après de fortes pluies (11,7, 12,2 et 7,6 mm) ; l'alarme d'humidité du boîtier est active les 22 et 23 et la tension tombe à 4,0 V le 22. "
            "Les 6, 13 et 27/09 sont partiels (7, 7 et 3 h). Ces 6 jours sont exclus des statistiques journalières comparatives."),
          fig("01_disponibilite", 16.5)]
story += [PageBreak(), P("3. Résultats", "h1"), P("3.1 Température et humidité", "h2"), fig("02_temperature", 16.5), fig("03_cycle_journalier", 16.5),
          P(f"Le maximum du jour est atteint vers 14 h UTC (33,5 °C en moyenne). Température et humidité sont fortement anticorrélées "
            f"(r = {T['corr'].loc['temperature','humidity']:.2f}), signe d'une masse d'air à humidité absolue quasi constante (rosée 23,2 ± 1,1 °C). "
            f"Heures valides ≥ 35 °C : {ind['Heures valides ≥ 35 °C']} ; ≥ 33 °C : {ind['Heures valides ≥ 33 °C']}.")]
story += [P("3.2 Rayonnement", "h2"), fig("04_rayonnement", 16.5),
          P("Pic d'éclairement à 11 h UTC (≈ 91 klux) et UV moyen de 6,2. Éclairement et UV sont corrélés à r = 1,00 : l'indice UV semble dérivé de l'éclairement, "
            "ce ne sont pas deux mesures indépendantes.")]
story += [PageBreak(), P("3.3 Vent", "h2"),
          Table([[fig("05_rose_vents", 8.2), fig("06_vent", 8.8)]], colWidths=[8.5 * cm, 8.8 * cm]),
          P(f"Vent faible : moyenne 1,5 m/s, {ind['Pourcentage de calme (vent = 0)']} % de calme. Rafale maximale {ind['Rafale max (m/s)']} m/s ({ind['Rafale max (km/h)']} km/h). "
            f"Secteurs dominants : N et NNE ({ind['Part secteurs N+NNE']} des vents à eux deux), avec un second lobe secondaire ESE et ONO. Régime de mousson : pas de direction exclusive.")]
story += [P("3.4 Pluie", "h2"), fig("07_pluie", 16.5),
          P("Le capteur renvoie un <b>compteur cumulatif</b> (334,3 → 392,9 mm). Les incréments sont convertis en pluie par intervalle ; aucune remise à zéro sur le mois. "
            "Les plus gros incréments (13,0 mm le 10/09, 12,2 mm le 22/09, 11,7 mm le 20/09) tombent à la <b>première mesure après un silence de nuit</b> : "
            "ils agrègent la pluie de plusieurs heures. La pluie fine vue en direct (0,1–0,3 mm/min) survient les 17 (matin), 19, 23, 24 et 29/09 (après-midi). "
            "<b>Le cumul mensuel de 58,6 mm est donc un minimum fiable, mais pas une répartition horaire.</b>")]
story += [PageBreak(), P("3.5 Pression et point de rosée", "h2"), fig("08_pression_rosee", 16.5),
          P(f"Pression moyenne 980,3 hPa (±1,9) avec une marée barométrique diurne nette : maximum vers {p_peak[0]} h UTC (982,5 hPa) et minimum vers {p_peak[1]} h (978,2 hPa). "
            "Pas de baisse de pression marquée sur le mois (minimum 974,9 hPa).")]
story += [P("3.6 Corrélations", "h2"),
          Table([[fig("09_correlations", 8.3), P("Lecture : <b>température ↔ humidité −0,93</b> (cohérence physique), <b>température ↔ pression −0,60</b> (marée barométrique diurne), "
            "<b>éclairement ↔ UV 1,00</b> (capteurs non indépendants), vent peu lié aux autres variables. Aucune corrélation suspecte qui trahirait un défaut de capteur.", "b")]],
            colWidths=[8.6 * cm, 8.6 * cm])]
story += [P("3.7 Alimentation et réseau", "h2"), fig("10_batterie_signal", 16.5),
          P("La tension de la batterie varie de 4,0 à 5,3 V : 4,4 V à 06 h, 5,3 V (pleine charge) de 09 h à 17 h — cohérent avec une recharge solaire. "
            "Le niveau de charge annoncé (<i>battery_level</i>) reste figé à 67 %, il n'est pas exploitable : utiliser la tension. Signal Wi-Fi moyen −80 dBm (min −99), correct mais limite.")]
story += [PageBreak(), P("4. Statistiques descriptives (fenêtre diurne, mesures nettoyées)", "h1"),
          tbl(S, [3.6 * cm, 1.2 * cm, 1.4 * cm] + [1.6 * cm] * 7), Spacer(1, .3 * cm),
          P("5. Synthèse journalière (jours exploitables)", "h1"),
          tbl(Df.reset_index().rename(columns={"index": "Jour"}).assign(Jour=lambda x: pd.to_datetime(x["Jour"]).dt.strftime("%d/%m"))[
              ["Jour", "heures_valides", "T_min", "T_max", "T_moy", "HR_moy", "vent_moy_ms", "rafale_max_ms", "UV_max", "pluie_mm"]]
              .rename(columns={"heures_valides": "h valides", "T_min": "Tmin", "T_max": "Tmax", "T_moy": "Tmoy", "HR_moy": "HR %", "vent_moy_ms": "Vent", "rafale_max_ms": "Rafale", "pluie_mm": "Pluie mm"}),
              [1.6 * cm] + [1.7 * cm] * 9), P("Tmin = minimum de la fenêtre diurne, pas minimum de la nuit.", "s")]
story += [PageBreak(), P("6. Recommandations", "h1"),
          P("<b>1. Comprendre pourquoi la station se tait la nuit</b> (priorité). Piste la plus probable : mise en veille faute de charge (tension 4,0–4,5 V le matin) ; "
            "à confirmer sur place ou avec le support Shelly. Tant que ce n'est pas corrigé, aucune donnée nocturne n'existera (T min, point de rosée de nuit, pluie datée)."),
          P("<b>2. Dédoublonner à la collecte.</b> Dans <i>shelly_cl</i>, 61,5 % des lignes de cette station sont des répétitions. Ne stocker une ligne que si <i>measured_at</i> change "
            "réduirait fortement le volume (voir le suivi de volume de la base) et simplifierait toutes les analyses."),
          P("<b>3. Dater la pluie.</b> Pour exploiter le cumul, exiger une transmission continue ; sinon, ne publier que des totaux par jour ou par nuit, jamais des intensités."),
          P("<b>4. Surveiller la batterie par la tension</b> et poser une alerte sous 4,3 V ; ne plus utiliser <i>battery_level</i>."),
          P("<b>5. Valider les valeurs</b> (température, pluie 58,6 mm) face à une station de référence de la région avant de les utiliser dans un rapport officiel."),
          P("7. Méthode et limites", "h1"),
          P("Dédoublonnage sur <i>measured_at</i> ; contrôles de bornes physiques ; filtre de Hampel (médiane glissante de 11 mesures, seuils : température 3 °C, humidité 15 %, pression 2 hPa, vent 4 m/s) ; "
            "détection des séquences figées (≥ 120 mesures) ; cohérence rafale/vent et rosée/température. Heure valide : ≥ 30 mesures sur 60. Jour exploitable : ≥ 8 h valides. "
            "Direction du vent moyennée vectoriellement (pondérée par la vitesse). Rose des vents : secteurs de 22,5°, vent ≥ 0,5 m/s. "
            "Limites : échantillon diurne uniquement, un seul mois, une seule station, pas de référence externe. Les scripts (<i>analyse.py</i>, <i>build_outputs.py</i>) reproduisent l'ensemble à l'identique.", "b")]

def foot(c, d):
    c.saveState(); c.setFont("Helvetica", 7.5); c.setFillColor(colors.HexColor("#777777"))
    c.drawString(2 * cm, 1.1 * cm, "Station météo Donsin — septembre 2026"); c.drawRightString(A4[0] - 2 * cm, 1.1 * cm, f"Page {d.page}"); c.restoreState()
SimpleDocTemplate(f"{OUT}/Meteo_Donsin_2026-09_rapport.pdf", pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                  title="Station météo Donsin — septembre 2026", author="").build(story, onFirstPage=foot, onLaterPages=foot)
shutil.copytree(OUT, DESK, dirs_exist_ok=True); print("OK", os.listdir(DESK))
